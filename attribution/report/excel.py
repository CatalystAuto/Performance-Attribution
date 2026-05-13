"""Write the run output: 3-sheet wide-format workbook + parquet sidecar.

Sheet layout (rows = tickers, columns = dates):
  Weights     — benchmark weight per (ticker, date)
  Returns     — daily return per (ticker, date), plus J803TR and JSAPY rows
  Calculation — contribution per (ticker, date), plus PORT/IDX/DIFF summary rows
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

def _pivot_wide(
    df: pd.DataFrame,
    *,
    index_col: str,
    value_col: str,
    dates: list[pd.Timestamp],
) -> pd.DataFrame:
    """Pivot long → wide, reindexing to exactly `dates` (fill_value=0)."""
    if df.empty:
        wide = pd.DataFrame(index=pd.Index([], name=index_col))
    else:
        wide = df.pivot_table(
            index=index_col,
            columns="date",
            values=value_col,
            aggfunc="first",
            fill_value=0,
        )
    # Ensure all required dates are present, in order
    wide = wide.reindex(columns=dates, fill_value=0)
    wide.columns.name = None
    return wide


def _build_body(
    wide: pd.DataFrame,
    ticker_order: list[str],
    ticker_to_bb: dict[str, str],
    ticker_to_name: dict[str, str],
    dates: list[pd.Timestamp],
) -> pd.DataFrame:
    """Prepend Ticker and Name columns; reindex rows to ticker_order."""
    wide = wide.reindex(index=ticker_order, fill_value=0)
    date_strs = [d.strftime("%Y-%m-%d") for d in dates]
    result = pd.DataFrame(index=range(len(ticker_order)))
    result["Ticker"] = [ticker_to_bb.get(t, t) for t in ticker_order]
    result["Name"] = [ticker_to_name.get(t, "") for t in ticker_order]
    for d, ds in zip(dates, date_strs):
        result[ds] = wide[d].values if d in wide.columns else 0
    return result


def write(
    out_xlsx: str | Path,
    *,
    weights: pd.DataFrame,           # long: date, ticker, benchmark_weight, weight_source
    returns: pd.DataFrame,           # long: date, ticker, daily_return
    contributions: pd.DataFrame,     # long: date, ticker, benchmark_weight, daily_return, contribution, weight_source
    daily: pd.DataFrame,             # date, benchmark_return_computed, benchmark_return_published, ...
    ignored_indices: pd.DataFrame,   # long: date, ticker, value — JSAPY etc.
    ticker_metadata: pd.DataFrame,   # ticker, name
) -> None:
    out_xlsx = Path(out_xlsx)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)

    # Determine universe and ordering
    universe = set(weights["ticker"].unique()) | set(contributions["ticker"].unique())
    # Use ticker_metadata order; only keep tickers present in universe
    ordered_tickers = [
        t for t in ticker_metadata["ticker"].tolist() if t in universe
    ]
    # Any ticker in universe not in ticker_metadata gets appended (fallback)
    for t in sorted(universe - set(ordered_tickers)):
        ordered_tickers.append(t)

    ticker_to_name: dict[str, str] = dict(
        zip(ticker_metadata["ticker"], ticker_metadata["name"])
    )
    ticker_to_bb: dict[str, str] = {t: f"{t} SJ" for t in universe}

    # Sort dates ascending — use weights as the canonical date set because
    # weights only exist for the requested month's trading dates, whereas the
    # returns frame also includes the base date (2026-02-28) and any other
    # dates present in the source xlsx.
    all_dates = sorted(
        pd.Series(weights["date"].dropna().unique()).tolist()
    )
    all_dates = [pd.Timestamp(d) for d in all_dates]
    date_strs = [d.strftime("%Y-%m-%d") for d in all_dates]

    # ── Weights wide pivot ──────────────────────────────────────────────────
    weights_wide = _pivot_wide(weights, index_col="ticker", value_col="benchmark_weight", dates=all_dates)
    weights_body = _build_body(weights_wide, ordered_tickers, ticker_to_bb, ticker_to_name, all_dates)

    # ── Returns wide pivot ──────────────────────────────────────────────────
    returns_wide = _pivot_wide(returns, index_col="ticker", value_col="daily_return", dates=all_dates)
    returns_body = _build_body(returns_wide, ordered_tickers, ticker_to_bb, ticker_to_name, all_dates)

    # Append J803TR row to Returns
    j803tr_row = {"Ticker": "J803TR", "Name": "J803TR Index"}
    daily_sorted = daily.sort_values("date")
    for d, ds in zip(all_dates, date_strs):
        match = daily_sorted[daily_sorted["date"] == d]
        j803tr_row[ds] = float(match["benchmark_return_published"].iloc[0]) if len(match) else 0
    returns_body = pd.concat(
        [returns_body, pd.DataFrame([j803tr_row])],
        ignore_index=True,
    )

    # Append JSAPY row to Returns
    jsapy_row = {"Ticker": "JSAPY", "Name": "JSAPY Index"}
    jsapy_data = ignored_indices[ignored_indices["ticker"] == "JSAPY Index"] if not ignored_indices.empty else pd.DataFrame()
    for d, ds in zip(all_dates, date_strs):
        if not jsapy_data.empty:
            match = jsapy_data[jsapy_data["date"] == d]
            jsapy_row[ds] = float(match["value"].iloc[0]) if len(match) else 0
        else:
            jsapy_row[ds] = 0
    returns_body = pd.concat(
        [returns_body, pd.DataFrame([jsapy_row])],
        ignore_index=True,
    )

    # ── Contributions wide pivot ────────────────────────────────────────────
    contribs_wide = _pivot_wide(contributions, index_col="ticker", value_col="contribution", dates=all_dates)
    contribs_body = _build_body(contribs_wide, ordered_tickers, ticker_to_bb, ticker_to_name, all_dates)

    # ── Build workbook with openpyxl ────────────────────────────────────────
    wb = openpyxl.Workbook()

    # Helper: write a DataFrame to a worksheet starting at a given row
    def _write_df_to_ws(ws, df: pd.DataFrame, start_row: int = 1) -> None:
        """Write df header + data to ws starting at start_row."""
        cols = list(df.columns)
        for c_idx, col in enumerate(cols, 1):
            ws.cell(start_row, c_idx).value = col
        for r_idx, row_data in enumerate(df.itertuples(index=False), start_row + 1):
            for c_idx, val in enumerate(row_data, 1):
                ws.cell(r_idx, c_idx).value = val if pd.notna(val) else None

    # ── Sheet 1: Weights ────────────────────────────────────────────────────
    ws_weights = wb.active
    ws_weights.title = "Weights"
    _write_df_to_ws(ws_weights, weights_body, start_row=1)

    # ── Sheet 2: Returns ────────────────────────────────────────────────────
    ws_returns = wb.create_sheet("Returns")
    _write_df_to_ws(ws_returns, returns_body, start_row=1)

    # ── Sheet 3: Calculation ────────────────────────────────────────────────
    ws_calc = wb.create_sheet("Calculation")

    # Row 1: header (Ticker, Name, date1, date2, ...)
    header = ["Ticker", "Name"] + date_strs
    for c_idx, h in enumerate(header, 1):
        ws_calc.cell(1, c_idx).value = h

    # Rows 2..1+N: Ticker + Name as literal labels (cols A/B); formula cells
    # for each date column reference the matching cell in Weights and Returns
    # so the user can audit by editing inputs.
    n_tickers = len(contribs_body)
    for r_offset, row_data in enumerate(contribs_body.itertuples(index=False)):
        r_idx = r_offset + 2
        ws_calc.cell(r_idx, 1).value = row_data[0]  # Ticker
        ws_calc.cell(r_idx, 2).value = row_data[1]  # Name
        for c_idx in range(3, len(header) + 1):
            col_letter = get_column_letter(c_idx)
            ws_calc.cell(r_idx, c_idx).value = (
                f"=Weights!{col_letter}{r_idx}*Returns!{col_letter}{r_idx}"
            )

    # Row after data: blank row (no values written → cells stay None)
    blank_row = n_tickers + 2  # first blank row

    # PORT / IDX / DIFF rows — all formulas.
    port_row_idx = blank_row + 1
    idx_row_idx = blank_row + 2
    diff_row_idx = blank_row + 3

    # PORT — label cols + SUM formula across the ticker rows in each date column
    ws_calc.cell(port_row_idx, 1).value = "PORT"
    ws_calc.cell(port_row_idx, 2).value = "Portfolio weighted return (SUM above)"
    for c_idx in range(3, len(header) + 1):
        col_letter = get_column_letter(c_idx)
        ws_calc.cell(port_row_idx, c_idx).value = (
            f"=SUM({col_letter}2:{col_letter}{n_tickers + 1})"
        )

    # IDX — label cols + reference the J803TR row in Returns. Returns has
    # the same N stock rows in the same order, then J803TR at row N+2.
    j803tr_returns_row = n_tickers + 2
    ws_calc.cell(idx_row_idx, 1).value = "IDX"
    ws_calc.cell(idx_row_idx, 2).value = "J803TR Index (from Returns)"
    for c_idx in range(3, len(header) + 1):
        col_letter = get_column_letter(c_idx)
        ws_calc.cell(idx_row_idx, c_idx).value = (
            f"=Returns!{col_letter}{j803tr_returns_row}"
        )

    # DIFF — label cols + (PORT - IDX) for each date column
    ws_calc.cell(diff_row_idx, 1).value = "DIFF"
    ws_calc.cell(diff_row_idx, 2).value = "Portfolio - J803TR"
    for c_idx in range(3, len(header) + 1):
        col_letter = get_column_letter(c_idx)
        ws_calc.cell(diff_row_idx, c_idx).value = (
            f"={col_letter}{port_row_idx}-{col_letter}{idx_row_idx}"
        )

    # ── Parquet sidecar ─────────────────────────────────────────────────────
    sidecar = out_xlsx.with_name(out_xlsx.stem + "_contributions.parquet")
    contributions.to_parquet(sidecar, index=False)

    wb.save(out_xlsx)
