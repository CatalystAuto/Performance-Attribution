import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from attribution.adapters import xlsx_returns
from attribution.compute import benchmark as bm


def test_end_to_end_on_synthetic_month(tmp_path, fixture_dir, monkeypatch):
    """Run the CLI end-to-end against tiny synthetic inputs.

    Layout:
      tmp/
        returns.xlsx                  <- copy of tiny_returns.xlsx
        pdfs/                          <- empty (forces CSV fallback)
        csvs/2026-03/J803/J803_20260302.csv, J803_20260303.csv
        out/
    """
    import shutil

    pdf_dir = tmp_path / "pdfs"
    pdf_dir.mkdir()
    csv_root = tmp_path / "csvs"
    csv_dir = csv_root / "2026-03" / "J803"
    csv_dir.mkdir(parents=True)
    out_dir = tmp_path / "out"

    returns_path = tmp_path / "returns.xlsx"
    shutil.copy(fixture_dir / "tiny_returns.xlsx", returns_path)

    csv_body = (
        "J803,Closing Weights\n"
        "AAA,0.4\n"
        "BBB,0.3\n"
        "CCC,0.3\n"
        "\n"
        "J803,Opening Weights\n"
        "AAA,0.4\n"
        "BBB,0.3\n"
        "CCC,0.3\n"
    )
    (csv_dir / "J803_20260302.csv").write_text(csv_body)
    (csv_dir / "J803_20260303.csv").write_text(csv_body)

    result = subprocess.run(
        [
            sys.executable, "-m", "attribution", "run",
            "--month", "2026-03",
            "--benchmark", "J803",
            "--pdf-dir", str(pdf_dir),
            "--csv-dir", str(csv_dir),
            "--returns", str(returns_path),
            "--out", str(out_dir),
        ],
        capture_output=True, text=True,
    )

    assert result.returncode in (0, 2), result.stderr

    out_xlsx = out_dir / "benchmark_attribution_2026-03.xlsx"
    assert out_xlsx.exists(), f"missing output: stderr={result.stderr}"
    wb = openpyxl.load_workbook(out_xlsx)
    assert "Weights" in wb.sheetnames
    assert "Returns" in wb.sheetnames
    assert "Calculation" in wb.sheetnames
